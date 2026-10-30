#!/usr/bin/env python3
"""Digital input alarms on the Alarm setup tab (docs/MODBUS_STANDARD_BLOCK.md section 7) against
build/netsim: 24 slots (alarm 30 + slot) with input, alarm when, delay, latch, severity and label;
Save writes only what changed (parameters 340-344, labels 350-373) and the page waits until it
reads back; the label then names the alarm in /api/alarms/config, in the Modbus map's
description for the slot's discrete input, and in the Excel handover file's Codes sheet; a bad
delay is refused on the page; LOCAL mode disables the editor. Firmware without the slots (no
parameter 340) keeps the panel hidden, and the test only checks that.
Needs Python Playwright, openpyxl and a Chromium."""
import asyncio, glob, io, json, os, ssl, subprocess, sys, tempfile, time, urllib.request
import openpyxl
from playwright.async_api import async_playwright
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CHROME = os.environ.get("CHROME") or (sorted(glob.glob("/opt/pw-browsers/chromium-*/chrome-linux*/chrome")) or [None])[-1]
MB, WEB = 15950, 18950
URL = f"https://127.0.0.1:{WEB}/"
CTX = ssl._create_unverified_context()
passed = failed = 0


def check(cond, what):
    global passed, failed
    print(("  ok   " if cond else "  FAIL ") + what)
    if cond: passed += 1
    else: failed += 1


def get(path):
    with urllib.request.urlopen(URL + path, context=CTX, timeout=5) as f:
        return json.loads(f.read())


def has_dia():
    for _ in range(50):
        try: return any(d["id"] == 340 for d in get("api/params"))
        except OSError: time.sleep(0.2)
    return False


async def run(say, real):
    async with async_playwright() as p:
        b = await p.chromium.launch(executable_path=CHROME)
        ctx = await b.new_context(viewport={"width": 1300, "height": 1000}, ignore_https_errors=True, bypass_csp=True, accept_downloads=True)
        pg = await ctx.new_page()
        errs = []
        pg.on("pageerror", lambda e: errs.append(str(e)))
        for _ in range(100):
            try: await pg.goto(URL); break
            except Exception: await asyncio.sleep(0.1)
        await pg.wait_for_timeout(1000)
        await pg.select_option("#lvl", "admin"); await pg.fill("#pw", "2222"); await pg.click("#loginbtn")
        await pg.wait_for_selector('#tabbar [data-tab="alarmcfg"]', state="visible", timeout=10000)
        await pg.click('#tabbar [data-tab="alarmcfg"]')
        await pg.wait_for_timeout(1500)
        if not real:
            check(not await pg.is_visible("#diasec"), "firmware without the slots: no Digital input alarms panel")
            check(errs == [], f"no script errors {errs}")
            await b.close(); return
        await pg.wait_for_selector("#diasec", state="visible", timeout=5000)
        check(len(await pg.query_selector_all("#diatab tbody tr")) == 24, "24 slots")
        local = await pg.is_disabled("#diasave")
        say("set 1 1"); await pg.wait_for_timeout(2500)              # REMOTE, unless the slots are any-mode
        await pg.click('#tabbar [data-tab="alarmcfg"]'); await pg.wait_for_timeout(800)
        check(await pg.is_enabled("#diasave"), f"editable at admin (LOCAL disabled: {local})")
        row = "#diatab tbody tr:nth-child(3)"                        # slot 2 = alarm 32
        await pg.select_option(f"{row} [data-f=input]", "5")
        await pg.select_option(f"{row} [data-f=sev]", "3")
        await pg.fill(f"{row} [data-f=delay]", "9999")
        await pg.click("#diasave")
        await pg.wait_for_function("document.getElementById('diaerr').textContent.length > 0", timeout=5000)
        check("delay of alarm 32" in await pg.inner_text("#diaerr"), "a bad delay is refused on the page: " + await pg.inner_text("#diaerr"))
        await pg.fill(f"{row} [data-f=delay]", "5")
        await pg.check(f"{row} [data-f=latch]")
        await pg.fill(f"{row} [data-f=label]", "Door open")
        await pg.evaluate("document.getElementById('diaerr').textContent = ''")
        await pg.click("#diasave")
        await pg.wait_for_function("document.getElementById('diaerr').textContent.length > 0", timeout=20000)
        msg = await pg.inner_text("#diaerr")
        check(msg.startswith("Saved (5 changes)"), f"only the changes are written ({msg})")
        v = get("api/values?ids=340,342,343,344,352")["v"]
        check(v["340"][2] == 5 and v["342"][2] == 5 and bool(v["343"][2]) and v["344"][2] == 3 and v["352"] == "Door open", "settings read back")
        al = next(a for a in get("api/alarms/config")["alarms"] if a["id"] == 32)
        check(al["text"] == "Door open", f"the label names alarm 32 ({al['text']!r})")
        try:
            mp = get("api/modbus/map")
            e = next((x for x in mp["entries"] if x["table"] == "discrete" and x["addr"] == 7002), None)
            check(e is not None and e.get("description") == "Door open", f"Modbus map: discrete input 7002 is 'Door open' ({e and e.get('description')!r})")
        except urllib.error.HTTPError:
            check(False, "Modbus map available")
        # the handover file: Codes lists alarm 32 by its label, Standard values carries 7002's description
        await pg.click('#tabbar [data-tab="service"]')
        await pg.wait_for_selector("#mbxlsx", state="visible", timeout=10000)
        async with pg.expect_download(timeout=15000) as dl:
            await pg.click("#mbxlsx")
        wb = openpyxl.load_workbook(io.BytesIO(open(await (await dl.value).path(), "rb").read()))
        codes = {r[1].value: r[2].value for r in wb["Codes"].iter_rows(min_row=2) if r[0].value == "Alarm ID"}
        check(codes.get(32) == "Door open", f"Excel Codes: alarm 32 is {codes.get(32)!r}")
        std = [[c.value for c in r] for r in wb["Standard values"].iter_rows()]
        check(any(r[1] == 7002 and r[10] == "Door open" for r in std[1:]), "Excel Standard values: discrete input 7002 described 'Door open'")
        # the slot's live state: input 5 is off, so the alarm stays normal
        await pg.click('#tabbar [data-tab="alarmcfg"]')
        await pg.wait_for_function("document.querySelector('#diatab tbody tr:nth-child(3) .badge').textContent.length > 0", timeout=8000)
        check(await pg.inner_text(f"{row} .badge") in ("normal", "active, not acknowledged", "input unreliable"), "slot status: " + await pg.inner_text(f"{row} .badge"))
        await pg.click("#logoutbtn"); await pg.wait_for_timeout(1500)
        check(not await pg.is_visible("#diasec"), "logged out: panel hidden")
        check(errs == [], f"no script errors {errs}")
        await b.close()


def main():
    tmp = tempfile.mkdtemp()
    unit = subprocess.Popen([os.path.join(ROOT, "build", "netsim"), str(MB), str(WEB), os.path.join(tmp, "f.bin"), os.path.join(tmp, "n.bin")],
                            stdin=subprocess.PIPE, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, text=True)
    def say(line): unit.stdin.write(line + "\n"); unit.stdin.flush()
    try:
        real = has_dia()
        print("digital input alarms:", "the firmware's own" if real else "firmware without the slots: hidden-panel check only")
        asyncio.run(run(say, real))
    finally:
        unit.terminate()
        try: unit.wait(timeout=10)
        except subprocess.TimeoutExpired: unit.kill()
    print(f"\nDI ALARMS WEB: {passed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
